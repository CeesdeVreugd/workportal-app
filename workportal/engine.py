"""Rekenkern voor calculatie en na-calculatie (los van de webpagina's, zodat
dezelfde berekening overal gelijk is: scherm, PDF, Excel en dashboard)."""
import io
import json
import re

from openpyxl import load_workbook

from .db import query
from .util import to_float, get_setting

HOUR_UNITS = {"uur", "uren", "u", "hr", "hrs", "hour", "hours"}


# ---------------------------------------------------------------- calculatie

def load_calc(cid):
    c = query("SELECT * FROM calculations WHERE id = ?", (cid,), one=True)
    if not c:
        return None
    data = dict(c)
    data["worktypes"] = [w for w in (c["worktypes"] or "").split(",") if w]
    try:
        data["answers"] = json.loads(c["answers"] or "{}")
    except ValueError:
        data["answers"] = {}
    parts = []
    for p in query("SELECT * FROM calc_parts WHERE calc_id = ? ORDER BY sort, id", (cid,)):
        kd = [{"d": k["description"], "p": k["unit_price"], "q": k["qty"]}
              for k in query("SELECT * FROM calc_koopdelen WHERE part_id = ? ORDER BY sort, id", (p["id"],))]
        secs = []
        for s in query("SELECT * FROM calc_sections WHERE part_id = ? ORDER BY sort, id", (p["id"],)):
            lines = [{"d": l["description"], "u": l["unit"] or "", "q": l["qty"], "p": l["price"], "f": l["factor"],
                      "kd": bool(l["is_koopdelen"])}
                     for l in query("SELECT * FROM calc_lines WHERE section_id = ? ORDER BY sort, id", (s["id"],))]
            secs.append({"category": s["category"], "title": s["title"], "erp_pos": s["erp_pos"], "lines": lines})
        parts.append({"name": p["name"], "koopdelen": kd, "sections": secs})
    data["parts"] = parts
    return data


def compute_calc(data):
    """Totalen per onderdeel/sectie en per ERP-positie."""
    by_pos = {}
    total = hours_total = base_total = 0.0
    parts_out = []
    for part in data["parts"]:
        kd_total = sum((to_float(k.get("p"), 0) or 0) * (to_float(k.get("q"), 0) or 0) for k in part.get("koopdelen", []))
        p_total = 0.0
        secs = []
        for s in part["sections"]:
            s_total = s_base = s_hours = 0.0
            for l in s["lines"]:
                q = to_float(l.get("q"), 0) or 0
                price = kd_total if l.get("kd") else (to_float(l.get("p"), 0) or 0)
                f = to_float(l.get("f"), 1)
                f = 1.0 if f is None else f
                s_base += q * price
                s_total += q * price * f
                if (l.get("u") or "").lower() in HOUR_UNITS:
                    s_hours += q
            secs.append({"category": s["category"], "title": s["title"], "erp_pos": s.get("erp_pos"),
                         "base": s_base, "total": s_total, "hours": s_hours})
            p_total += s_total
            if s.get("erp_pos"):
                pos = int(s["erp_pos"])
                agg = by_pos.setdefault(pos, {"total": 0.0, "hours": 0.0, "title": s["title"]})
                agg["total"] += s_total
                agg["hours"] += s_hours
            hours_total += s_hours
            base_total += s_base
        parts_out.append({"name": part["name"], "total": p_total, "koopdelen": kd_total, "sections": secs})
        total += p_total
    return {"total": total, "base": base_total, "hours": hours_total, "parts": parts_out, "by_pos": by_pos}


def calc_total(cid):
    data = load_calc(cid)
    return compute_calc(data)["total"] if data else 0.0


# ---------------------------------------------------------------- ERP-export (regie)

def _norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


COLS = {
    "itemid": "item_id", "itemomschrijving": "item_desc", "pos": "pos", "aantal": "qty", "regeltype": "line_type",
    "eenheid": "unit", "artikelnr": "article", "omschrijving": "description",
    "eenheidprijsink": "cost_unit", "prijsfactorink": "cost_factor", "kortingink": "cost_discount",
    "subtotaalink": "cost_subtotal", "marge": "margin",
    "eenheidprijsvkp": "sale_unit", "eeneheidprijsvkp": "sale_unit", "prijsfactorvkp": "sale_factor",
    "kortingvkp": "sale_discount", "subtotaalvkp": "sale_subtotal",
}
HEAD = {"ordernr": "order_no", "orderomschrijving": "order_desc", "ordertotaal": "order_total",
        "gefactureerdtotaal": "invoiced_total"}


def parse_erp_xlsx(data: bytes):
    """Leest de regie-Excel uit het ERP. Geeft (kop, regels) terug.
    Kolommen worden op naam herkend, zodat een extra kolom niets breekt."""
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    header, lines = {}, []
    i = 0
    while i < len(rows):
        norm = [_norm(v) for v in rows[i]]
        if "ordernr" in norm and not header:
            nxt = rows[i + 1] if i + 1 < len(rows) else []
            for ci, key in enumerate(norm):
                if key in HEAD and ci < len(nxt):
                    header[HEAD[key]] = nxt[ci]
            i += 2
            continue
        if "itemid" in norm and "pos" in norm:
            colmap = {ci: COLS[k] for ci, k in enumerate(norm) if k in COLS}
            for r in rows[i + 1:]:
                if not r or all(v is None or str(v).strip() == "" for v in r):
                    continue
                rec = {}
                for ci, field in colmap.items():
                    rec[field] = r[ci] if ci < len(r) else None
                if rec.get("pos") in (None, ""):
                    continue
                for f in ("qty", "cost_unit", "cost_factor", "cost_discount", "cost_subtotal", "margin",
                          "sale_unit", "sale_factor", "sale_discount", "sale_subtotal"):
                    rec[f] = to_float(rec.get(f), 0.0)
                rec["pos"] = int(to_float(rec.get("pos"), 0))
                rec["line_type"] = int(to_float(rec.get("line_type"), 0))
                for f in ("item_id", "item_desc", "unit", "article", "description"):
                    v = rec.get(f)
                    if isinstance(v, float) and v.is_integer():
                        v = int(v)
                    rec[f] = None if v is None else str(v).strip()
                lines.append(rec)
            break
        i += 1
    wb.close()
    if not lines:
        raise ValueError("Geen regels gevonden. Is dit de regie-export uit het ERP (kolommen ItemID, Pos, Aantal ...)?")
    header["order_no"] = None if header.get("order_no") is None else str(
        int(header["order_no"]) if isinstance(header["order_no"], float) else header["order_no"]).strip()
    header["order_total"] = to_float(header.get("order_total"))
    header["invoiced_total"] = to_float(header.get("invoiced_total"))
    return header, lines


def _is_header_line(l):
    return l["line_type"] == 1 or (l["pos"] % 100 == 0 and not l["cost_subtotal"] and not l["sale_subtotal"])


def _is_hours(l):
    return (l.get("unit") or "").lower() in HOUR_UNITS or l.get("line_type") == 4


def compute_nacalc(n, lines, calc_data=None):
    """n: dict nacalcs-rij; lines: lijst dicts. Geeft (samenvatting, items)."""
    try:
        opts = json.loads(n.get("no_divide") or "{}")
    except ValueError:
        opts = {}
    if isinstance(opts, list):
        opts = {"items": opts}
    no_div = set(int(x) for x in opts.get("items", []))
    divide_order = opts.get("order", True)
    N = to_float(n.get("divide_by"), 1) or 1

    items = {}
    order = []
    for l in lines:
        key = l.get("item_id") or str(l["pos"] // 100 * 100)
        if key not in items:
            items[key] = {"key": key, "item_id": l.get("item_id"), "desc": l.get("item_desc") or l.get("description") or "",
                          "pos": None, "cost": 0.0, "sale": 0.0, "hours": 0.0, "lines": []}
            order.append(key)
        it = items[key]
        if _is_header_line(l):
            it["pos"] = l["pos"]
            if l.get("item_desc"):
                it["desc"] = l["item_desc"]
            continue
        if it["pos"] is None:
            it["pos"] = l["pos"] // 100 * 100
        it["cost"] += l.get("cost_subtotal") or 0
        it["sale"] += l.get("sale_subtotal") or 0
        if _is_hours(l):
            it["hours"] += l.get("qty") or 0
        it["lines"].append(l)

    calc = compute_calc(calc_data) if calc_data else None
    out = []
    for key in order:
        it = items[key]
        divided = N != 1 and it["pos"] not in no_div
        if divided:
            it["cost"] /= N
            it["sale"] /= N
            it["hours"] /= N
        it["divided"] = divided
        it["calc"] = None
        it["calc_hours"] = None
        if calc:
            c = calc["by_pos"].get(it["pos"])
            if c:
                it["calc"] = c["total"]
                it["calc_hours"] = c["hours"]
        it["diff"] = (it["calc"] - it["sale"]) if it["calc"] is not None else None
        it["margin"] = it["sale"] - it["cost"]
        if it["calc"] is not None:
            if it["sale"] <= it["calc"] + 0.005:
                it["status"], it["label"] = "ok", "Binnen budget"
            elif it["sale"] <= it["calc"] * 1.10:
                it["status"], it["label"] = "warn", "Let op"
            else:
                it["status"], it["label"] = "bad", "Overschreden"
        else:
            it["status"], it["label"] = ("info", "Niet gecalculeerd") if calc else ("", "")
        out.append(it)

    if calc:
        known = {it["pos"] for it in out}
        for pos, c in sorted(calc["by_pos"].items()):
            if pos not in known and (c["total"] or c["hours"]):
                out.append({"key": f"calc-{pos}", "item_id": None, "desc": c["title"] + " (alleen gecalculeerd)", "pos": pos,
                            "cost": 0.0, "sale": 0.0, "hours": 0.0, "lines": [], "divided": False, "calc": c["total"],
                            "calc_hours": c["hours"], "diff": c["total"], "margin": 0.0, "status": "ok",
                            "label": "Nog geen kosten"})
    out.sort(key=lambda x: (x["pos"] is None, x["pos"] or 0))

    order_total = to_float(n.get("order_total"))
    if not order_total and calc_data and calc_data.get("agreed_price"):
        order_total = to_float(calc_data.get("agreed_price"))
    order_total = order_total or 0.0
    if N != 1 and divide_order:
        order_total /= N
    cost = sum(i["cost"] for i in out)
    sale = sum(i["sale"] for i in out)
    hours = sum(i["hours"] for i in out)
    result = order_total - cost
    margin_pct = (result / order_total * 100) if order_total else 0.0
    extra = to_float(get_setting("extra_margin_pct"), 5) or 0
    if order_total <= 0:
        status, label = "info", "Geen orderbedrag"
    elif order_total < cost:
        status, label = "bad", "Verlies"
    elif order_total < sale - 0.005:
        status, label = "warn", "Onder minimale marge"
    elif order_total < sale * (1 + extra / 100):
        status, label = "ok", "Minimale marge gehaald"
    else:
        status, label = "okplus", "Extra bovenop marge"
    summary = {
        "order_total": order_total, "cost": cost, "sale": sale, "hours": hours, "result": result,
        "margin_pct": margin_pct, "status": status, "label": label, "min_margin": sale - cost,
        "above_min": order_total - sale, "divide_by": N, "divide_order": divide_order, "no_divide": sorted(no_div),
        "extra_pct": extra,
    }
    if calc:
        summary["calc_total"] = calc["total"]
        summary["calc_hours"] = calc["hours"]
        summary["diff"] = calc["total"] - sale
    return summary, out
