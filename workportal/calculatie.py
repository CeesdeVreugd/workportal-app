import io
import json

from flask import Blueprint, render_template, request, redirect, url_for, flash, abort, g, jsonify, Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .db import query, execute, get_db
from .engine import load_calc, compute_calc
from .permissions import require, can, LEZEN, BEWERKEN, BEHEER
from .seed import CATEGORIES, CATEGORY_TITLES, WORKTYPES
from .klanten import customer_options, ask_snelstart
from .util import now_iso, audit, to_float, to_int, next_number, safe_filename

bp = Blueprint("calculatie", __name__, url_prefix="/calculatie")

STATUSES = [("concept", "Concept"), ("verstuurd", "Verstuurd"), ("akkoord", "Akkoord"), ("verloren", "Niet doorgegaan")]
UNITS = ["uur", "stuk", "post", "rit", "km", "m", "m2", "kg", "set"]


def _f(name):
    v = request.form.get(name)
    return v.strip() if v and v.strip() else None


def _templates():
    out = {}
    for r in query("SELECT * FROM calc_templates"):
        try:
            out[r["worktype"]] = json.loads(r["data"])
        except ValueError:
            out[r["worktype"]] = {}
    return out


def build_structure(worktypes, part_names, include):
    """Bouwt onderdelen/secties/regels op uit de sjablonen van de gekozen werktypes."""
    tpls = _templates()
    parts = []
    pos = 100
    for pname in part_names:
        secs = []
        for cat, title, _ in CATEGORIES:
            if not include.get(cat, True):
                continue
            seen, lines = set(), []
            for wt in worktypes:
                for l in tpls.get(wt, {}).get(cat, []):
                    key = (l.get("d") or "").strip().lower()
                    if key and key not in seen:
                        seen.add(key)
                        lines.append({"d": l["d"], "u": l.get("u", ""), "q": 0, "p": l.get("p", 0), "f": l.get("f", 1), "kd": False})
            if cat == "inkoop" and lines:
                lines.append({"d": "Koopdelen (zie koopdelen specificatie)", "u": "post", "q": 1, "p": 0, "f": 1.25, "kd": True})
            if not lines:
                continue
            secs.append({"category": cat, "title": title, "erp_pos": pos, "lines": lines})
            pos += 100
        parts.append({"name": pname, "koopdelen": [], "sections": secs})
    return parts


def save_structure(cid, parts):
    db = get_db()
    db.execute("DELETE FROM calc_parts WHERE calc_id = ?", (cid,))
    for pi, part in enumerate(parts):
        cur = db.execute("INSERT INTO calc_parts (calc_id, sort, name) VALUES (?,?,?)",
                         (cid, pi, (part.get("name") or f"Onderdeel {pi + 1}")[:120]))
        pid = cur.lastrowid
        for ki, k in enumerate(part.get("koopdelen") or []):
            if not (k.get("d") or "").strip() and not to_float(k.get("p"), 0):
                continue
            db.execute("INSERT INTO calc_koopdelen (part_id, sort, description, unit_price, qty) VALUES (?,?,?,?,?)",
                       (pid, ki, (k.get("d") or "").strip()[:300], to_float(k.get("p"), 0) or 0, to_float(k.get("q"), 0) or 0))
        for si, s in enumerate(part.get("sections") or []):
            cur = db.execute("INSERT INTO calc_sections (part_id, sort, category, title, erp_pos) VALUES (?,?,?,?,?)",
                             (pid, si, s.get("category") or "overig", (s.get("title") or "Sectie")[:200],
                              to_int(s.get("erp_pos"))))
            sid = cur.lastrowid
            for li, l in enumerate(s.get("lines") or []):
                desc = (l.get("d") or "").strip()
                if not desc and not to_float(l.get("q"), 0):
                    continue
                f = to_float(l.get("f"), 1)
                db.execute("INSERT INTO calc_lines (section_id, sort, description, unit, qty, price, factor, is_koopdelen)"
                           " VALUES (?,?,?,?,?,?,?,?)",
                           (sid, li, desc[:300] or "–", (l.get("u") or "")[:20], to_float(l.get("q"), 0) or 0,
                            to_float(l.get("p"), 0) or 0, 1.0 if f is None else f, 1 if l.get("kd") else 0))
    db.commit()


def _calc_or_404(cid):
    data = load_calc(cid)
    if not data:
        abort(404)
    return data


@bp.route("/")
@require("calculatie", LEZEN)
def index():
    view = request.args.get("view", "open")
    q = (request.args.get("q") or "").strip()
    where, params = [], []
    if view == "open":
        where.append("c.status IN ('concept','verstuurd')")
    elif view in dict(STATUSES):
        where.append("c.status = ?")
        params.append(view)
    if q:
        where.append("(c.number LIKE ? OR c.title LIKE ? OR IFNULL(c.customer_text,'') LIKE ? OR IFNULL(cu.name,'') LIKE ?)")
        params += [f"%{q}%"] * 4
    sql = ("SELECT c.*, cu.name AS customer, p.number AS project_no, u.name AS author FROM calculations c"
           " LEFT JOIN customers cu ON cu.id = c.customer_id LEFT JOIN projects p ON p.id = c.project_id"
           " LEFT JOIN users u ON u.id = c.created_by")
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY c.updated_at DESC LIMIT 300"
    calcs = []
    for c in query(sql, params):
        d = dict(c)
        d["total"] = compute_calc(load_calc(c["id"]))["total"]
        calcs.append(d)
    return render_template("calculatie/index.html", calcs=calcs, view=view, q=q, STATUSES=STATUSES)


@bp.route("/nieuw", methods=["GET", "POST"])
@require("calculatie", BEWERKEN)
def new():
    customers = customer_options(to_int(request.form.get("customer_id")))
    projects = query("SELECT id, number, name, customer_id FROM projects ORDER BY number DESC")
    if request.method == "POST":
        worktypes = [w for w in request.form.getlist("worktypes") if w in dict(WORKTYPES)]
        title = _f("title")
        if not title or not worktypes:
            flash("Vul een omschrijving in en kies minimaal één type werk.", "error")
            return render_template("calculatie/new.html", form=request.form, customers=customers, projects=projects,
                                   WORKTYPES=WORKTYPES)
        n_parts = max(1, min(10, to_int(request.form.get("n_parts"), 1)))
        names = [(_f(f"part_{i}") or (title if n_parts == 1 else f"Onderdeel {i + 1}")) for i in range(n_parts)]
        include = {
            "montage": request.form.get("montage") == "ja",
            "besturing": request.form.get("besturing") == "ja",
            "transport": request.form.get("transport") != "nee",
        }
        answers = {k: request.form.get(k) for k in ("montage", "besturing", "transport", "material", "distance_km")}
        answers["expected"] = _f("expected")
        cust_id = to_int(request.form.get("customer_id"))
        now = now_iso()
        cid = execute(
            "INSERT INTO calculations (number, title, customer_id, customer_text, project_id, worktypes, answers, quantity,"
            " agreed_price, status, notes, created_by, created_at, updated_at) VALUES (?,?,?,?,?,?,?,?,?, 'concept', ?, ?, ?, ?)",
            (next_number("calculations", "C", 3), title, cust_id, _f("customer_text"), to_int(request.form.get("project_id")),
             ",".join(worktypes), json.dumps(answers), max(1, to_int(request.form.get("quantity"), 1)),
             to_float(request.form.get("expected")), _f("notes"), g.user["id"], now, now))
        save_structure(cid, build_structure(worktypes, names, include))
        audit("aangemaakt", "calc", cid, title)
        flash("Calculatie opgebouwd uit de sjablonen. Vul de aantallen en prijzen in.", "ok")
        target = url_for("calculatie.edit", cid=cid)
        return ask_snelstart(cust_id, target) or redirect(target)
    return render_template("calculatie/new.html", form={"n_parts": 1, "quantity": 1, "transport": "ja"},
                           customers=customers, projects=projects, WORKTYPES=WORKTYPES)


@bp.route("/<int:cid>")
@require("calculatie", LEZEN)
def edit(cid):
    data = _calc_or_404(cid)
    customers = customer_options(data["customer_id"])
    projects = query("SELECT id, number, name, customer_id FROM projects ORDER BY number DESC")
    nacalcs = query("SELECT id, imported_at, order_no FROM nacalcs WHERE calc_id = ? ORDER BY imported_at DESC", (cid,))
    readonly = not can("calculatie", BEWERKEN)
    return render_template("calculatie/edit.html", c=data, customers=customers, projects=projects, nacalcs=nacalcs,
                           STATUSES=STATUSES, UNITS=UNITS, CATEGORIES=CATEGORIES, WORKTYPES=dict(WORKTYPES),
                           readonly=readonly)


@bp.route("/<int:cid>/opslaan", methods=["POST"])
@require("calculatie", BEWERKEN)
def save(cid):
    _calc_or_404(cid)
    js = request.get_json(silent=True) or {}
    parts = js.get("parts")
    if not isinstance(parts, list) or not parts:
        return jsonify({"error": "Geen onderdelen ontvangen"}), 400
    status = js.get("status") if js.get("status") in dict(STATUSES) else "concept"
    execute("UPDATE calculations SET title=?, customer_id=?, customer_text=?, project_id=?, quantity=?, agreed_price=?,"
            " status=?, notes=?, updated_at=? WHERE id=?",
            ((js.get("title") or "Calculatie")[:200], to_int(js.get("customer_id")), (js.get("customer_text") or None),
             to_int(js.get("project_id")), max(1, to_int(js.get("quantity"), 1)), to_float(js.get("agreed_price")),
             status, js.get("notes") or None, now_iso(), cid))
    save_structure(cid, parts)
    totals = compute_calc(load_calc(cid))
    return jsonify({"ok": True, "total": totals["total"], "saved_at": now_iso()})


@bp.route("/<int:cid>/kopie", methods=["POST"])
@require("calculatie", BEWERKEN)
def duplicate(cid):
    data = _calc_or_404(cid)
    now = now_iso()
    new_id = execute(
        "INSERT INTO calculations (number, title, customer_id, customer_text, project_id, worktypes, answers, quantity,"
        " agreed_price, status, notes, created_by, created_at, updated_at) VALUES (?,?,?,?,?,?,?,?,?, 'concept', ?, ?, ?, ?)",
        (next_number("calculations", "C", 3), data["title"] + " (kopie)", data["customer_id"], data["customer_text"],
         data["project_id"], ",".join(data["worktypes"]), json.dumps(data["answers"]), data["quantity"], data["agreed_price"],
         data["notes"], g.user["id"], now, now))
    save_structure(new_id, data["parts"])
    audit("gekopieerd", "calc", new_id, data["number"])
    flash("Kopie gemaakt.", "ok")
    return redirect(url_for("calculatie.edit", cid=new_id))


@bp.route("/<int:cid>/verwijderen", methods=["POST"])
@require("calculatie", BEHEER)
def delete(cid):
    data = _calc_or_404(cid)
    execute("DELETE FROM calculations WHERE id = ?", (cid,))
    audit("verwijderd", "calc", cid, data["number"])
    flash(f"Calculatie {data['number']} verwijderd.", "ok")
    return redirect(url_for("calculatie.index"))


@bp.route("/<int:cid>/excel")
@require("calculatie", LEZEN)
def excel(cid):
    data = _calc_or_404(cid)
    totals = compute_calc(data)
    cust = data["customer_text"]
    if data["customer_id"]:
        row = query("SELECT name FROM customers WHERE id = ?", (data["customer_id"],), one=True)
        cust = row["name"] if row else cust
    wb = Workbook()
    wb.remove(wb.active)
    bold = Font(bold=True)
    blue = Font(bold=True, color="0A0A96", size=12)
    head_fill = PatternFill("solid", fgColor="F3F5FA")
    thin = Side(style="thin", color="E1E5F0")
    eur = '"€" #,##0.00'
    part_refs = []
    kd_sheet = wb.create_sheet("Koopdelen specificatie")
    kd_sheet["A1"] = "Calculatie koopdelen specificatie"
    kd_sheet["A1"].font = blue
    kd_row = 3
    kd_totals = {}
    for pi, part in enumerate(data["parts"]):
        kd_sheet.cell(kd_row, 1, f"Koopdelen {part['name']}").font = bold
        kd_row += 1
        for ci, h in enumerate(["Benaming", "Stuksprijs", "Aantal", "Totaal"], 1):
            c = kd_sheet.cell(kd_row, ci, h)
            c.font = bold
            c.fill = head_fill
        kd_row += 1
        first = kd_row
        for k in part["koopdelen"]:
            kd_sheet.cell(kd_row, 1, k["d"])
            kd_sheet.cell(kd_row, 2, k["p"]).number_format = eur
            kd_sheet.cell(kd_row, 3, k["q"])
            kd_sheet.cell(kd_row, 4, f"=B{kd_row}*C{kd_row}").number_format = eur
            kd_row += 1
        kd_sheet.cell(kd_row, 1, "Totaal").font = bold
        kd_sheet.cell(kd_row, 4, f"=SUM(D{first}:D{max(first, kd_row - 1)})" if part["koopdelen"] else 0).number_format = eur
        kd_totals[pi] = f"'Koopdelen specificatie'!$D${kd_row}"
        kd_row += 3
    kd_sheet.column_dimensions["A"].width = 46
    for col in "BCD":
        kd_sheet.column_dimensions[col].width = 14

    for pi, part in enumerate(data["parts"]):
        ws = wb.create_sheet(safe_filename(part["name"])[:28] or f"Tabblad {pi + 1}", pi)
        ws["A1"] = "Calculatie"
        ws["A1"].font = Font(bold=True, color="0A0A96", size=16)
        ws["A3"], ws["B3"] = "Klant:", cust or ""
        ws["A4"], ws["B4"] = "Omschrijving", f"{data['title']} – {part['name']}"
        ws["A5"], ws["B5"] = "Calculatienr.", data["number"]
        for ci, h in enumerate(["", "Aantal", "Eenheid", "Prijs per stuk", "marge+ risico:", "Totaal", "Inclusief marge", "totaal"], 1):
            c = ws.cell(7, ci, h)
            c.font = bold
            c.fill = head_fill
        r = 9
        sec_refs = []
        for s in part["sections"]:
            ws.cell(r, 1, f"{s['title']}  (pos. {s.get('erp_pos') or '-'})").font = blue
            r += 2
            first = r
            for l in s["lines"]:
                ws.cell(r, 1, l["d"])
                ws.cell(r, 2, l["q"])
                ws.cell(r, 3, l["u"])
                if l.get("kd"):
                    ws.cell(r, 4, "=" + kd_totals[pi]).number_format = eur
                else:
                    ws.cell(r, 4, l["p"]).number_format = eur
                ws.cell(r, 5, l["f"])
                ws.cell(r, 6, f"=B{r}*D{r}").number_format = eur
                ws.cell(r, 7, f"=B{r}*D{r}*E{r}").number_format = eur
                r += 1
            last = r - 1
            ws.cell(r, 1, "totaal").font = bold
            ws.cell(r, 8, f"=SUM(G{first}:G{max(first, last)})").number_format = eur
            ws.cell(r, 8).font = bold
            sec_refs.append((s["title"], f"H{r}"))
            r += 3
        ws.cell(r, 1, "Totalen").font = blue
        r += 2
        first = r
        for title, ref in sec_refs:
            ws.cell(r, 1, title)
            ws.cell(r, 8, f"={ref}").number_format = eur
            r += 1
        ws.cell(r + 1, 1, "Totaal onderdeel").font = bold
        ws.cell(r + 1, 8, f"=SUM(H{first}:H{r - 1})").number_format = eur
        ws.cell(r + 1, 8).font = bold
        part_refs.append((ws.title, f"H{r + 1}"))
        ws.column_dimensions["A"].width = 48
        for col, w in zip("BCDEFGH", (9, 9, 14, 13, 14, 16, 16)):
            ws.column_dimensions[col].width = w

    tot = wb.create_sheet("Totalen", len(data["parts"]))
    tot["A1"] = "Calculatie totalen"
    tot["A1"].font = Font(bold=True, color="0A0A96", size=16)
    tot["A3"], tot["C3"] = "Klant:", cust or ""
    tot["A4"], tot["C4"] = "Projectnaam:", data["title"]
    headers = ["Nr.", "Omschrijving", "Totalen"] + [t for _, t, _ in CATEGORIES]
    for ci, h in enumerate(headers, 1):
        c = tot.cell(6, ci, h)
        c.font = bold
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True)
    r = 7
    for pi, (part, calc_part) in enumerate(zip(data["parts"], totals["parts"])):
        tot.cell(r, 1, pi + 1)
        tot.cell(r, 2, part["name"])
        tot.cell(r, 3, f"='{part_refs[pi][0]}'!{part_refs[pi][1]}").number_format = eur
        for ci, (cat, _, _) in enumerate(CATEGORIES, 4):
            val = sum(s["total"] for s in calc_part["sections"] if s["category"] == cat)
            tot.cell(r, ci, round(val, 2)).number_format = eur
        r += 1
    tot.cell(r + 1, 2, "Totaalprijs").font = bold
    tot.cell(r + 1, 3, f"=SUM(C7:C{r - 1})").number_format = eur
    tot.cell(r + 2, 2, "Prijsafspraak").font = bold
    tot.cell(r + 2, 3, data["agreed_price"] or None).number_format = eur
    if data["quantity"] and data["quantity"] > 1:
        tot.cell(r + 3, 2, f"Prijs per stuk ({data['quantity']} st.)")
        tot.cell(r + 3, 3, f"=C{r + 1}/{data['quantity']}").number_format = eur
    tot.column_dimensions["B"].width = 30
    for i in range(3, 10):
        tot.column_dimensions[get_column_letter(i)].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    fname = safe_filename(f"{data['number']} - Calculatie {data['title']}.xlsx")
    return Response(buf.getvalue(), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{fname}"'})
