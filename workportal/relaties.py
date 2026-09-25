"""Import van relaties (klanten en leveranciers) uit een Excel-export van Komdex.

Werking:
- Elke relatie uit Komdex heeft een vast ID. Relaties in WorkPortal die al aan dat ID
  gekoppeld zijn, worden bijgewerkt (ook een gewijzigde naam).
- Alleen velden die in Komdex zijn veranderd sinds de vorige import worden overgenomen,
  zodat een bewuste keuze voor de gegevens uit de app niet steeds wordt overschreven.
- Een relatie uit Komdex zonder koppeling die lijkt op een relatie die in de app is
  aangemaakt (zelfde naam of e-mailadres), wordt één keer voorgelegd: app-versie houden,
  Excel-versie gebruiken of als aparte relatie toevoegen. Daarna is het ID gekoppeld en
  komt de vraag niet terug.
"""
import io
import json
import os
import re
import secrets

from flask import current_app
from openpyxl import load_workbook

from .db import get_db
from .util import now_iso

RELATION_TYPES = [
    ("klant", "Klant"),
    ("leverancier", "Leverancier"),
    ("beide", "Klant en leverancier"),
    ("instelling", "Instelling"),
    ("prospect", "Prospect"),
]
TYPE_LABELS = dict(RELATION_TYPES)

HEADERS = {
    "id": "komdex_id", "relatieid": "komdex_id",
    "verkortenaam": "short_name", "zoeknaam": "short_name",
    "naam": "name", "bedrijfsnaam": "name",
    "adres": "street", "straat": "street",
    "huisnr": "house_no", "huisnummer": "house_no",
    "pc": "postcode", "postcode": "postcode",
    "plaats": "city", "woonplaats": "city",
    "tel": "phone", "telefoon": "phone", "telefoonnummer": "phone",
    "email": "email", "emailadres": "email",
    "website": "website", "www": "website",
    "relatietype": "type",
    "relatiegroep": "group", "groep": "group", "branche": "group",
}

# Velden die uit Komdex komen (het klantnummer SnelStart en notities horen bij de app)
FIELDS = ["short_name", "name", "address", "postcode", "city", "phone", "email", "website",
          "relation_type", "relation_group", "active"]
FIELD_LABELS = {"short_name": "Verkorte naam", "name": "Naam", "address": "Adres", "postcode": "Postcode",
                "city": "Plaats", "phone": "Telefoon", "email": "E-mail", "website": "Website",
                "relation_type": "Type", "relation_group": "Groep", "active": "Actief"}

LEGAL = r"\b(b\.?\s?v\.?|v\.?\s?o\.?\s?f\.?|n\.?\s?v\.?|holding|groep|group|nederland|the|de|het)\b"


def _norm(v):
    return re.sub(r"[^a-z0-9]", "", str(v or "").lower())


def name_key(v):
    s = str(v or "").lower().replace("&", " en ")
    s = re.sub(LEGAL, " ", s)
    return re.sub(r"[^a-z0-9]", "", s)


def normalize_type(raw):
    """Relatietype uit Komdex -> (vaste sleutel, actief)."""
    t = _norm(raw)
    if not t or t == "leeg":
        return None, True
    if "vervallen" in t:
        return None, False
    if ("klant" in t or "deb" in t) and ("lever" in t or "cred" in t):
        return "beide", True
    if "lever" in t or "cred" in t:
        return "leverancier", True
    if "klant" in t or "deb" in t:
        return "klant", True
    if "instelling" in t:
        return "instelling", True
    if "prospect" in t:
        return "prospect", True
    return None, True


def _is_type(v):
    t = _norm(v)
    return bool(t) and any(k in t for k in ("klant", "lever", "debiteur", "crediteur", "instelling", "vervallen",
                                            "prospect", "leeg"))


def _text(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def _phone(v):
    t = _text(v)
    if t and t.startswith("="):
        t = "+" + t[1:].lstrip("+")
    return t


def parse_relations(data: bytes):
    """Leest de Komdex-export. Geeft (records, waarschuwingen)."""
    # data_only=False: telefoonnummers als "+31 ..." staan in Komdex-exports soms als formule (=31345-569641)
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    hdr_idx = None
    for i, r in enumerate(rows[:10]):
        keys = [_norm(v) for v in r]
        if "naam" in keys or "bedrijfsnaam" in keys:
            hdr_idx = i
            break
    if hdr_idx is None:
        raise ValueError("Geen kolom 'Naam' gevonden. Is dit de relatie-export uit Komdex?")
    colmap = {}
    for ci, v in enumerate(rows[hdr_idx]):
        key = HEADERS.get(_norm(v))
        if key and key not in colmap.values():
            colmap[ci] = key
    if "komdex_id" not in colmap.values():
        raise ValueError("Geen kolom 'ID' gevonden. Het Komdex-ID is nodig om relaties te kunnen koppelen.")
    inv = {v: k for k, v in colmap.items()}
    street_col = inv.get("street", 10 ** 6)
    type_col = inv.get("type")
    out, warnings, seen = [], [], set()
    for rn, r in enumerate(rows[hdr_idx + 1:], start=hdr_idx + 2):
        if not r or all(v is None or str(v).strip() == "" for v in r):
            continue
        # Verschoven regel herstellen: het relatietype staat dan een paar kolommen verder
        shift = 0
        if type_col is not None and not _is_type(r[type_col] if type_col < len(r) else None):
            for k in range(1, 4):
                if type_col + k < len(r) and _is_type(r[type_col + k]):
                    shift = k
                    break
        rec = {}
        for ci, key in colmap.items():
            src = ci + shift if ci >= street_col else ci
            rec[key] = r[src] if src < len(r) else None
        kid = _text(rec.get("komdex_id"))
        name = _text(rec.get("name"))
        if not kid:
            warnings.append(f"Regel {rn}: geen ID, overgeslagen")
            continue
        if not name or _norm(name) == "leeg":
            warnings.append(f"Regel {rn} (ID {kid}): geen naam, overgeslagen")
            continue
        if kid in seen:
            warnings.append(f"Regel {rn}: ID {kid} komt dubbel voor, overgeslagen")
            continue
        seen.add(kid)
        street = _text(rec.get("street"))
        house = _text(rec.get("house_no"))
        address = street
        if street and house and not street.endswith(house):
            address = f"{street} {house}"
        elif not street and house:
            address = house
        rtype, active = normalize_type(rec.get("type"))
        group = _text(rec.get("group"))
        if group and _norm(group) == "leeg":
            group = None
        out.append({
            "komdex_id": kid,
            "short_name": _text(rec.get("short_name")),
            "name": name,
            "address": address,
            "postcode": _text(rec.get("postcode")),
            "city": _text(rec.get("city")),
            "phone": _phone(rec.get("phone")),
            "email": _text(rec.get("email")),
            "website": _text(rec.get("website")),
            "relation_type": rtype,
            "relation_group": group,
            "active": 1 if active else 0,
            "_row": rn,
        })
        if shift:
            warnings.append(f"Regel {rn} ({name}): kolommen verschoven in de export, automatisch rechtgezet")
    return out, warnings


# ---------------------------------------------------------------- tijdelijke opslag tussen stap 1 en 2

def _dir():
    d = os.path.join(current_app.config["DATA_DIR"], "imports")
    os.makedirs(d, exist_ok=True)
    return d


def store(records, warnings, filename):
    token = secrets.token_hex(12)
    with open(os.path.join(_dir(), f"{token}.json"), "w") as fh:
        json.dump({"records": records, "warnings": warnings, "filename": filename, "at": now_iso()}, fh)
    return token


def load(token):
    if not re.fullmatch(r"[0-9a-f]{24}", token or ""):
        return None
    path = os.path.join(_dir(), f"{token}.json")
    if not os.path.exists(path):
        return None
    with open(path) as fh:
        return json.load(fh)


def discard(token):
    try:
        os.remove(os.path.join(_dir(), f"{token}.json"))
    except OSError:
        pass


# ---------------------------------------------------------------- plan en uitvoeren

def _snapshot(row):
    try:
        return json.loads(row["komdex_snapshot"] or "{}")
    except ValueError:
        return {}


def _changes(row, rec):
    """Velden die in Komdex zijn gewijzigd t.o.v. de vorige import (of t.o.v. de app als er nog niets is)."""
    snap = _snapshot(row)
    out = {}
    for f in FIELDS:
        new = rec.get(f)
        base = snap[f] if f in snap else row[f]
        if (new or None) != (base or None) and (new or None) != (row[f] or None):
            out[f] = (row[f], new)
    return out


def plan(records):
    db = get_db()
    linked = {r["komdex_id"]: r for r in db.execute("SELECT * FROM customers WHERE komdex_id IS NOT NULL")}
    loose = db.execute("SELECT * FROM customers WHERE komdex_id IS NULL").fetchall()
    by_name, by_email = {}, {}
    for r in loose:
        for k in {name_key(r["name"]), name_key(r["short_name"])} - {""}:
            by_name.setdefault(k, []).append(r)
        if r["email"]:
            by_email.setdefault(r["email"].strip().lower(), []).append(r)
    result = {"new": [], "updates": [], "unchanged": 0, "conflicts": []}
    for rec in records:
        row = linked.get(rec["komdex_id"])
        if row is not None:
            ch = _changes(row, rec)
            if ch:
                result["updates"].append({"id": row["id"], "rec": rec, "changes": ch})
            else:
                result["unchanged"] += 1
            continue
        cands = {}
        for k in {name_key(rec["name"]), name_key(rec["short_name"])} - {""}:
            for r in by_name.get(k, []):
                cands[r["id"]] = r
        if rec.get("email"):
            for r in by_email.get(rec["email"].strip().lower(), []):
                cands[r["id"]] = r
        if cands:
            result["conflicts"].append({"rec": rec, "candidates": [dict(c) for c in cands.values()]})
        else:
            result["new"].append(rec)
    return result


def _insert(db, rec, now):
    vals = [rec.get(f) for f in FIELDS]
    snap = json.dumps({f: rec.get(f) for f in FIELDS})
    db.execute(f"INSERT INTO customers (komdex_id, komdex_snapshot, {', '.join(FIELDS)}, created_at, updated_at)"
               f" VALUES (?, ?, {', '.join('?' * len(FIELDS))}, ?, ?)", [rec["komdex_id"], snap] + vals + [now, now])


def apply(records, decisions):
    """decisions: {komdex_id: 'app:<id>' | 'excel:<id>' | 'apart'}"""
    db = get_db()
    now = now_iso()
    p = plan(records)
    stats = {"nieuw": 0, "bijgewerkt": 0, "ongewijzigd": p["unchanged"], "gekoppeld": 0, "namen": [], "dubbel": []}
    for rec in p["new"]:
        _insert(db, rec, now)
        stats["nieuw"] += 1
    for u in p["updates"]:
        sets = {f: new for f, (_, new) in u["changes"].items()}
        if "name" in sets:
            stats["namen"].append((u["changes"]["name"][0], sets["name"]))
        snap = json.dumps({f: u["rec"].get(f) for f in FIELDS})
        db.execute(f"UPDATE customers SET {', '.join(f + ' = ?' for f in sets)}, komdex_snapshot = ?, updated_at = ?"
                   f" WHERE id = ?", list(sets.values()) + [snap, now, u["id"]])
        stats["bijgewerkt"] += 1
    used = set()
    for c in p["conflicts"]:
        rec = c["rec"]
        choice = decisions.get(rec["komdex_id"], "")
        kind, _, cid = choice.partition(":")
        cid = int(cid) if cid.isdigit() else None
        valid = {cand["id"] for cand in c["candidates"]}
        snap = json.dumps({f: rec.get(f) for f in FIELDS})
        if kind in ("app", "excel") and cid in valid and cid not in used:
            used.add(cid)
            if kind == "excel":
                db.execute(f"UPDATE customers SET komdex_id = ?, komdex_snapshot = ?, "
                           f"{', '.join(f + ' = ?' for f in FIELDS)}, updated_at = ? WHERE id = ?",
                           [rec["komdex_id"], snap] + [rec.get(f) for f in FIELDS] + [now, cid])
            else:
                db.execute("UPDATE customers SET komdex_id = ?, komdex_snapshot = ?, updated_at = ? WHERE id = ?",
                           (rec["komdex_id"], snap, now, cid))
            stats["gekoppeld"] += 1
        else:
            if kind in ("app", "excel") and cid in used:
                stats["dubbel"].append(rec["name"])
            _insert(db, rec, now)
            stats["nieuw"] += 1
    db.commit()
    return stats
